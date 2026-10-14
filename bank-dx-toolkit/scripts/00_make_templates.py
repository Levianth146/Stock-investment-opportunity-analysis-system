"""Tạo các file mẫu: sources.csv, coding_sheet.xlsx, manual_audit.xlsx, dti_inputs.csv"""
import csv, yaml
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
from common import *

banks = load_yaml("banks.yaml")
kw = load_yaml("keywords.yaml")
B = list(banks["banks"].keys()); years = banks["years"]; cyears = banks["coding_years"]

# 1) sources.csv: danh sách tài liệu cần tải (điền URL thủ công)
p = CONF / "sources.csv"
if not p.exists():
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["bank", "year", "doc_type", "url", "notes"])
        for b in B:
            for y in years:
                for t in ("bctn", "bctc", "bcptbv"):
                    w.writerow([b, y, t, "", ""])
    print("tạo", p)

# 2) coding_sheet.xlsx: bảng chấm kép
# Idempotent: nếu đã có file thì GIỮ mọi điểm đã chấm, chỉ thêm các dòng (tiêu chí x ngân hàng x năm) còn thiếu rồi sắp lại theo thứ tự chuẩn.
OLD = {}
if (OUT / "coding_sheet.xlsx").exists():
    from openpyxl import load_workbook
    _ws = load_workbook(OUT / "coding_sheet.xlsx").active
    for r in _ws.iter_rows(min_row=2, values_only=True):
        if r[0]: OLD[(r[0], r[2], int(r[3]))] = list(r)
wb = Workbook(); ws = wb.active; ws.title = "coding"
hdr = ["criterion", "criterion_name", "bank", "year", "coder1", "coder2", "final", "not_disclosed",
       "evidence_file", "evidence_page", "evidence_quote", "notes"]
ws.append(hdr)
for c in hdr: ws.cell(1, hdr.index(c) + 1).font = Font(bold=True); ws.cell(1, hdr.index(c) + 1).fill = PatternFill("solid", fgColor="D9D9D9")
for code, v in kw.items():
    for b in B:
        for y in cyears:
            ws.append(OLD.get((code, b, y)) or [code, v["name"], b, y, None, None, None, None, None, None, None, None])
dv = DataValidation(type="list", formula1='"0,1,2,3,4"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"E2:G{ws.max_row}")
dv2 = DataValidation(type="list", formula1='"0,1"', allow_blank=True); ws.add_data_validation(dv2); dv2.add(f"H2:H{ws.max_row}")
for col, wd in zip("ABCDEFGHIJKL", [9, 38, 10, 7, 8, 8, 8, 12, 28, 10, 60, 30]): ws.column_dimensions[col].width = wd
ws.freeze_panes = "E2"
wb.save(OUT / "coding_sheet.xlsx"); print("tạo", OUT / "coding_sheet.xlsx")

# 3) manual_audit.xlsx: kiểm kê 8 nền tảng số
if not (OUT / "manual_audit.xlsx").exists():
    PLAT = ["Mobile banking KHCN", "Ngân hàng số KHDN", "Mở TK/eKYC trực tuyến", "QR/ví/thanh toán nhanh",
            "Vay số trực tuyến", "Open API/cổng developer", "Chatbot/trợ lý ảo", "Xác thực sinh trắc học/Smart OTP"]
    wb = Workbook(); ws = wb.active; ws.title = "audit"
    ws.append(["bank", "platform_no", "platform", "achieved(0/1)", "screenshot_or_url", "audit_date", "tester", "notes"])
    for b in B:
        for i, pl in enumerate(PLAT, 1):
            ws.append([b, i, pl, None, None, None, None, None])
    dv3 = DataValidation(type="list", formula1='"0,1"', allow_blank=True); ws.add_data_validation(dv3); dv3.add(f"D2:D{ws.max_row}")
    for col, wd in zip("ABCDEFGH", [10, 11, 32, 14, 40, 12, 12, 30]): ws.column_dimensions[col].width = wd
    wb.save(OUT / "manual_audit.xlsx"); print("tạo", OUT / "manual_audit.xlsx")

else:
    print("giữ bảng kiểm kê đã có", OUT / "manual_audit.xlsx")

# 4) dti_inputs.csv: nơi gom số liệu đã XÁC NHẬN cho 11 chỉ số DTI + biến Lớp C
VARS = ["digital_txn_share", "digital_customers", "total_customers", "headcount", "tech_staff", "training_hours_per_head",
        "software_gross_cost", "total_assets", "opex", "op_profit_before_provision", "net_profit", "equity",
        "casa_deposits", "customer_deposits", "net_interest_income", "net_fee_income"]
p = OUT / "dti_inputs.csv"
if not p.exists():
    with open(p, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f); w.writerow(["bank", "year", "variable", "value", "unit", "source_file", "page", "confirmed_by", "note"])
        for b in B:
            for y in years:
                for v in VARS:
                    w.writerow([b, y, v, "", "", "", "", "", ""])
    print("tạo", p)
