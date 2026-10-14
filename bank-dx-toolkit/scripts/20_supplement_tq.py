"""Bổ sung kiểm kê T-Q2 từ bằng chứng công khai đã đọc; ghi ngoại lệ T-Q1 VCB.

Chạy từ scripts/: python 20_supplement_tq.py. Không đăng nhập/kiểm thử giao dịch.
Các lựa chọn nguồn được đối chiếu ngày 2026-10-14, không chấm bằng từ khóa.
"""
import csv
import json
import re
from openpyxl import load_workbook
from common import OUT, DATA

AUDIT_DATE = "2026-10-14"
METHOD = "public_document_inventory_v1"
# Mã nền tảng giữ nguyên checklist 00; mỗi bộ gồm năm, trang PDF, cụm định vị.
PDF_EVIDENCE = {
    "agribank": {1: (2025, 18, "Agribank Plus"), 2: (2025, 16, "Agribank eBanking"),
                 4: (2025, 16, "mã QR"), 5: (2025, 18, "cho vay trực tuyến"),
                 6: (2025, 16, "OpenAPI"), 7: (2025, 45, "AI Chatbot")},
    "mb": {1: (2025, 88, "App MBBank"), 2: (2025, 88, "BIZ MBBank"),
           3: (2025, 88, "account opening"), 4: (2024, 8, "QR code"),
           5: (2025, 88, "entirely online"), 6: (2025, 19, "API connections"),
           7: (2025, 88, "AI-powered chatbots"), 8: (2025, 88, "online biometric")},
    "tcb": {1: (2025, 53, "Techcombank Mobile"), 2: (2025, 53, "business banking application"),
            4: (2025, 32, "QR"), 5: (2025, 53, "loan origination"),
            6: (2024, 29, "API capabilities"), 7: (2025, 54, "Chatbot Triage"),
            8: (2025, 62, "biometric authentication")},
    "bidv": {1: (2025, 50, "SmartBanking"), 2: (2025, 45, "BIDV Direct"),
             5: (2025, 7, "Salary Advance Loan"), 6: (2024, 51, "BIDV Open API"),
             7: (2025, 43, "Chatbot solutions"), 8: (2024, 15, "Biometric transaction authentication")},
    "vcb": {1: (2025, 57, "VCB Digibank"), 2: (2025, 57, "VCB DigiBiz"),
            4: (2025, 58, "thanh toán QR"), 5: (2025, 21, "VCB Online Lending"),
            7: (2025, 57, "Chatbot"), 8: (2025, 58, "SoftOTP")},
}
WEB_EVIDENCE = {
    ("agribank", 3): (
        "https://www.agribank.com.vn/vn/ve-agribank/tin-tuc/dtl?current=true&urile=wcm:path:/agbank/ve-agribank/tin-tuc/cac-tin-khac/tuoi-tre-agribank/diem-giao-dich-xanh-agribank-chuong-trinh-hanh-dong-gan-lien-xu-the-phat-trien-ben-vung",
        "Tuần lễ Văn hóa Agribank 2025; mô tả mở tài khoản và dịch vụ số mọi lúc, mọi nơi."),
    ("agribank", 8): (
        "https://www.agribank.com.vn/vn/ve-agribank/tin-tuc-su-kien/tin-ve-agribank/hoat-dong-agribank/agribank-dong-hanh-cung-cong-uoc-ha-noi-kien-tao-khong-gian-so-an-toan-phat-trien-ben-vung",
        "Bài về Công ước Hà Nội 25–26/10/2025 xác nhận các giải pháp xác thực Soft-OTP, token và khuôn mặt."),
    ("tcb", 3): (
        "https://techcombank.com/khach-hang-ca-nhan/ngan-hang-truc-tuyen/ngan-hang-so/techcombank-mobile/huong-dan-su-dung",
        "Hướng dẫn sản phẩm có mục mở tài khoản thanh toán bằng eKYC hoàn toàn trực tuyến; ảnh chụp tài liệu tại ngày kiểm kê."),
    ("bidv", 3): (
        "https://bidv.com.vn/bidv/tin-tuc/tin-ve-bidv/bidv-smartbanking-x-lap-ky-luc-viet-nam",
        "Bài ra mắt SmartBanking X tháng 7/2025 xác nhận tiện ích eKYC."),
    ("bidv", 4): (
        "https://bidv.com.vn/bidv/tin-tuc/tin-ve-bidv/bidv-smartbanking-x-lap-ky-luc-viet-nam",
        "Bài SmartBanking X năm 2025 xác nhận chuyển tiền QR."),
    ("vcb", 3): (
        "https://vietcombank.com.vn/vi-VN/Trang-thong-tin-dien-tu/Articles/2025/01/23/20250123_Don-loc-dau-xuan",
        "Chương trình ngày 03/02/2025 có điều kiện mở tài khoản trực tuyến và kích hoạt Digibank bằng eKYC."),
    ("vcb", 6): (
        "https://www.vietcombank.com.vn/en/Trang-thong-tin-dien-tu/Articles/2025/09/10/VCB-vinh-du-dat-giai-Euromoney-2025",
        "Bài công bố tháng 9/2025 xác nhận kết nối H2H/API phục vụ khách hàng doanh nghiệp."),
}


def main():
    wb = load_workbook(OUT / "manual_audit.xlsx")
    ws = wb.active
    headers = [c.value for c in ws[1]]
    extra = ["source_file", "page", "evidence_quote", "evidence_year", "scoring_year", "audit_method"]
    for name in extra:
        if name not in headers:
            headers.append(name)
            ws.cell(1, len(headers), name)
    evidence = []
    seen = set()
    for row in ws.iter_rows(min_row=2):
        bank, no = row[0].value, row[1].value
        key = bank, no
        if key in seen:
            raise ValueError(f"Trùng ô kiểm kê: {key}")
        seen.add(key)
        values = dict(zip(headers, [c.value for c in row]))
        if no in PDF_EVIDENCE.get(bank, {}):
            year, page, phrase = PDF_EVIDENCE[bank][no]
            records = [json.loads(line) for line in
                       (DATA / "text" / f"{bank}_{year}_bctn.jsonl").read_text(encoding="utf-8-sig").splitlines()]
            text = next(r["text"] for r in records if r["page"] == page)
            # Chỉ kiểm tra cụm định vị còn tồn tại; quyết định đạt nằm ở danh sách đã đọc ở trên.
            compact = re.sub(r"\s+", " ", text)
            match = re.search(re.escape(phrase), compact, re.I)
            if not match:
                raise ValueError(f"Không tìm thấy bằng chứng {bank}/{no}: {phrase}")
            quote = compact[max(0, match.start() - 120):match.end() + 380]
            source = f"data/raw/{bank}/{year}_bctn.pdf"
            values.update(source_file=source, page=page, evidence_quote=quote, evidence_year=year,
                          screenshot_or_url=f"{source}#page={page}")
        else:
            url, description = WEB_EVIDENCE[key]
            values.update(source_file=url, page=None, evidence_quote=description,
                          evidence_year=None if key == ("tcb", 3) else 2025, screenshot_or_url=url)
        values.update({"achieved(0/1)": 1, "audit_date": AUDIT_DATE,
                       "tester": "Codex — đối chiếu tài liệu công khai", "scoring_year": 2025,
                       "audit_method": METHOD,
                       "notes": "Có bằng chứng triển khai; không kiểm thử giao dịch. Kiểm kê tài liệu tại 2026-10-14, áp cho năm cuối 2025 theo bước 10."})
        for col, name in enumerate(headers, 1):
            ws.cell(row[0].row, col, values.get(name))
        evidence.append(values)
    if seen != {(b, i) for b in PDF_EVIDENCE for i in range(1, 9)}:
        raise ValueError("Checklist phải có đúng 40 ô thuộc 5 ngân hàng và 8 mã nền tảng")
    with (OUT / "platform_evidence.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(evidence)
    wb.save(OUT / "manual_audit.xlsx")

    # Không đánh đồng bản quyền/bằng sáng chế với phần mềm khi nguồn không tách riêng.
    path = OUT / "dti_inputs.csv"
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        fields = reader.fieldnames
        rows = list(reader)
    target = next(r for r in rows if (r["bank"], r["year"], r["variable"]) ==
                  ("vcb", "2025", "software_gross_cost"))
    if target["value"].strip():
        raise ValueError("T-Q1 VCB đã có số; cần đối chiếu trước khi thay đổi")
    target.update(source_file="data/raw/vcb/2025_bctc.pdf", page="49", confirmed_by="",
                  note="Đã đọc PDF 2026-10-14: thuyết minh 14, trang in 47 / PDF 49. Nguyên giá cuối năm: Bản quyền và bằng sáng chế 2.754.936 triệu VND; TSCĐ vô hình khác 15.543; tổng 5.284.812 gồm quyền sử dụng đất. Không tách riêng phần mềm. Giữ trống software_gross_cost, không lấy khoản gộp hoặc giá trị còn lại. Ảnh: out/tq_review/vcb_2025_bctc_p49.png.")
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    print("Đã bổ sung 40 ô T-Q2; T-Q1 VCB giữ trống vì chưa có nguyên giá riêng phần mềm.")


if __name__ == "__main__":
    main()
