"""Kiểm thử hồi quy offline: tải công khai, dữ liệu xác nhận, công thức và ngôn ngữ.
Chạy: python -B -m unittest discover -s tests -p test_regressions.py -v
Mọi tệp ghi thử đều nằm trong thư mục tạm, không sửa dữ liệu thật."""
import csv
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook
import requests

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import common
from public_web import PublicFetcher


def load_script(name):
    spec = importlib.util.spec_from_file_location(name.replace(".", "_"), ROOT / "scripts" / name)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


S = load_script("10_score.py")
D = load_script("01_download.py")


class TestDownload(unittest.TestCase):
    def parse(self, content):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sources.csv"
            path.write_text(content, encoding="utf-8-sig")
            return list(D.iter_source_rows(path))

    def test_extra_columns_are_rejected(self):
        self.assertEqual(self.parse("bank,year,doc_type,url,notes\nmb,2024,bctn,https://example.com,ok,extra\n"), [])

    def test_empty_extra_columns_do_not_crash(self):
        self.assertEqual(self.parse("bank,year,doc_type,url,notes\n,,,,,extra\n"), [])

    def test_missing_required_cell_is_rejected(self):
        self.assertEqual(self.parse("bank,year,doc_type,url\nmb,2024,bctn\n"), [])

    def test_blank_url_and_optional_report_remain_valid(self):
        rows = self.parse("bank,year,doc_type,url,notes\nmb,2024,bcptbv,,tích hợp trong BCTN\n")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["url"], "")

    def test_double_bom_and_header_spaces(self):
        rows = self.parse("\ufeffbank, year ,doc_type,url\nmb,2024,bctn,https://example.com\n")
        self.assertEqual(rows[0]["bank"], "mb")

    def test_unsafe_path_and_private_url_are_rejected(self):
        for bank, url in [("../mb", "https://example.com"), ("mb", "file:///private"), ("mb", "https://user:pass@example.com")]:
            with self.subTest(bank=bank, url=url):
                self.assertEqual(self.parse(f"bank,year,doc_type,url\n{bank},2024,bctn,{url}\n"), [])

    def test_retry_success_resets_status(self):
        fetcher = Mock()
        fetcher.get.side_effect = [requests.Timeout("thử lại"), Mock(content=b"%PDF-new")]
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "report.pdf"
            self.assertEqual(D.download_pdf(fetcher, "https://example.com", dest), "ok")
            self.assertEqual(dest.read_bytes(), b"%PDF-new")

    def test_bad_response_preserves_old_pdf(self):
        fetcher = Mock()
        fetcher.get.return_value = Mock(content=b"<html>captcha</html>")
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "report.pdf"
            dest.write_bytes(b"%PDF-old")
            self.assertTrue(D.download_pdf(fetcher, "https://example.com", dest).startswith("not_pdf"))
            self.assertEqual(dest.read_bytes(), b"%PDF-old")

    def test_failed_force_does_not_attribute_old_pdf_to_new_url(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "data/raw/mb/2024_bctn.pdf"
            dest.parent.mkdir(parents=True)
            dest.write_bytes(b"%PDF-old")
            (root / "sources.csv").write_text("bank,year,doc_type,url\nmb,2024,bctn,https://example.com/new.pdf\n", encoding="utf-8-sig")
            with patch.object(D, "CONF", root), patch.object(D, "DATA", root / "data"), patch.object(D, "ROOT", root), patch.object(D, "download_pdf", return_value="error: timeout"), patch.object(D.sys, "argv", ["01_download.py", "--force"]):
                D.main()
            with (root / "data/raw/manifest.csv").open(encoding="utf-8-sig") as f:
                row = next(csv.DictReader(f))
            self.assertEqual(row["file"], "")
            self.assertEqual(row["sha256"], "")
            self.assertEqual(dest.read_bytes(), b"%PDF-old")

    def test_import_does_not_download_or_write_manifest(self):
        with patch.object(PublicFetcher, "get", side_effect=AssertionError("Không được truy cập mạng")):
            with patch("builtins.open", side_effect=AssertionError("Không được mở manifest khi import")):
                load_script("01_download.py")


class TestPublicAccess(unittest.TestCase):
    def test_robots_failure_blocks_document(self):
        session = Mock(headers={})
        session.get.side_effect = requests.Timeout()
        fetcher = PublicFetcher("academic", session=session)
        with self.assertRaises(PermissionError):
            fetcher.get("https://example.com/report.pdf")
        self.assertEqual(session.get.call_count, 1)

    def test_robots_disallow_blocks_document(self):
        session = Mock(headers={})
        session.get.return_value = Mock(status_code=200, text="User-agent: *\nDisallow: /private")
        with self.assertRaises(PermissionError):
            PublicFetcher("academic", session=session).get("https://example.com/private/report.pdf")
        self.assertEqual(session.get.call_count, 1)

    def test_all_requests_include_delay_and_check_redirect_host(self):
        now = [0.0]
        timestamps = []
        responses = iter([
            Mock(status_code=404),
            Mock(status_code=302, headers={"Location": "https://other.example/report.pdf"}),
            Mock(status_code=404),
            Mock(status_code=200, content=b"%PDF"),
        ])
        session = Mock(headers={})
        def get(url, **kwargs):
            timestamps.append((url, now[0]))
            self.assertFalse(kwargs["allow_redirects"])
            return next(responses)
        session.get.side_effect = get
        with patch("public_web.time.monotonic", side_effect=lambda: now[0]), patch("public_web.time.sleep", side_effect=lambda dt: now.__setitem__(0, now[0] + dt)):
            fetcher = PublicFetcher("academic", delay=0, session=session)
            self.assertEqual(fetcher.get("https://example.com/report.pdf").status_code, 200)
        self.assertEqual([t for url, t in timestamps], [0, 3, 6, 9])
        self.assertEqual(timestamps[2][0], "https://other.example/robots.txt")


class TestScoreRegression(unittest.TestCase):
    def test_step0_threshold_zero_and_year_isolation(self):
        banks = ["a", "b", "c", "d", "e"]
        calc = Mock()
        # Một ngân hàng thiếu năm 2024: giữ; hai thiếu năm 2025: loại cả 5.
        calc.all.side_effect = lambda b, y: {"K-Q1": (None if b in (["e"] if y == 2024 else ["d", "e"]) else 0, "")}
        frame, gate = S.build_indicators(calc, banks, [2024, 2025], {"K-Q1": (0, 100, "higher_better")})
        old = frame[frame.year == 2024]
        new = frame[frame.year == 2025]
        self.assertEqual(old.included_in_score.sum(), 4)  # Số 0 là dữ liệu thật.
        self.assertTrue(new.z.isna().all())
        self.assertEqual(new.x.notna().sum(), 3)  # Giữ số gốc khi loại.
        self.assertEqual(new.z_before_step0.notna().sum(), 3)
        self.assertEqual(gate.n_missing.tolist(), [1, 2])

    def test_step0_excluded_pillar_uses_dbi(self):
        calc = Mock()
        calc.all.side_effect = lambda b, y: {"C-Q1": (10 if b == "a" else None, "")}
        frame, _ = S.build_indicators(calc, ["a", "b", "c"], [2025], {"C-Q1": (0, 100, "higher_better")})
        z = frame.loc[frame.bank == "a", "z"].iloc[0]
        scores = S.pillar_scores({"C1": 75, "C2": 75, "C3": 75, "C4": 75},
                                {"C-Q1": None if pd.isna(z) else z}, .8)
        self.assertEqual(scores["C"], (75, None, 75, 0))

    def test_latest_snapshot_only_and_last_store_row(self):
        df = pd.DataFrame([
            dict(bank="mb", store="apple", date="2026-10-01", rating=1, rating_count=100),
            dict(bank="mb", store="apple", date="2026-10-12", rating=2, rating_count=100),
            dict(bank="mb", store="apple", date="2026-10-12", rating=5, rating_count=100),
            dict(bank="mb", store="play", date="2026-10-12", rating=3, rating_count=300),
        ])
        with patch.object(S, "read_csv", return_value=df):
            self.assertEqual(S.load_app_rating("unused"), {"mb": (3.5, "2026-10-12")})

    def test_latest_invalid_snapshot_is_not_replaced_by_old_rating(self):
        df = pd.DataFrame([
            dict(bank="mb", store="apple", date="2026-10-01", rating=5, rating_count=100),
            dict(bank="mb", store="apple", date="2026-10-12", rating=0, rating_count=100),
        ])
        with patch.object(S, "read_csv", return_value=df):
            self.assertEqual(S.load_app_rating("unused"), {})

    def test_negative_opex_cir_and_toi(self):
        V = {("mb", 2025, "opex"): -20, ("mb", 2025, "op_profit_before_provision"): 80,
             ("mb", 2024, "headcount"): 5, ("mb", 2025, "headcount"): 15}
        calc = S.Calc(V, {}, {}, {}, 2021, 2025)
        self.assertEqual(calc.cir("mb", 2025), 20)
        self.assertEqual(calc.toi_per_head("mb", 2025), 10)

    def test_cagr_zero_base_is_missing(self):
        V = {("mb", 2021, "digital_customers"): 0, ("mb", 2025, "digital_customers"): 100}
        self.assertIsNone(S.Calc(V, {}, {}, {}, 2021, 2025).sq2("mb", 2025)[0])

    def test_normalization_nonfinite_values_remain_missing(self):
        for x in [np.nan, np.inf, -np.inf, None]:
            self.assertIsNone(S.normalize(x, 0, 100, "higher_better"))
        self.assertIsNone(S.normalize(50, 100, 0, "higher_better"))

    def test_audit_rejects_invalid_values_and_duplicate_platforms(self):
        for values, platforms in [([2] * 8, list(range(1, 9))), ([1] * 8, [1] * 8)]:
            df = pd.DataFrame({"bank": ["mb"] * 8, "platform_no": platforms, "achieved(0/1)": values})
            with patch.object(S.pd, "read_excel", return_value=df):
                self.assertEqual(S.load_platform("unused"), {})

    def test_audit_valid_zero_is_real_data(self):
        df = pd.DataFrame({"bank": ["mb"] * 8, "platform_no": range(1, 9), "achieved(0/1)": [0] * 8})
        with patch.object(S.pd, "read_excel", return_value=df):
            self.assertEqual(S.load_platform("unused"), {"mb": 0})

    def test_official_inputs_require_manual_confirmation_and_provenance(self):
        row = dict(bank="mb", year=2025, variable="digital_txn_share", value=80,
                   confirmed_by="nam", source_file="data/raw/mb/2025_bctn.pdf", page=2)
        for changes in [{"confirmed_by": "range-checked"}, {"confirmed_by": "fin-cand"}, {"source_file": ""}, {"page": np.nan}]:
            with patch.object(S, "read_csv", return_value=pd.DataFrame([{**row, **changes}])):
                with self.assertRaises(SystemExit):
                    S.load_inputs("unused")
        with patch.object(S, "read_csv", return_value=pd.DataFrame([row])):
            self.assertEqual(S.load_inputs("unused")[0][("mb", 2025, "digital_txn_share")], 80)

    def test_duplicate_inputs_are_rejected(self):
        row = dict(bank="mb", year=2025, variable="headcount", value=100)
        with patch.object(S, "read_csv", return_value=pd.DataFrame([row, row])):
            with self.assertRaises(SystemExit):
                S.load_inputs("unused", draft=True)

    def test_vietnamese_only_cannot_be_overridden(self):
        with patch.object(common, "ONLY_LANG", {"en"}):
            self.assertFalse(common.keep_lang("en", "vi"))
            self.assertFalse(common.keep_lang("vi", "vi"))
            self.assertTrue(common.keep_lang("en"))


class TestIsolatedPipeline(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="bankdx_regression_")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name) / "repo"
        for directory in ("scripts", "config"):
            shutil.copytree(ROOT / directory, self.root / directory, ignore=shutil.ignore_patterns("__pycache__"))
        self.run_script("00_make_templates.py")

    def run_script(self, name, *args):
        result = subprocess.run([sys.executable, "-B", name, *args], cwd=self.root / "scripts",
                                env=dict(os.environ, PYTHONUTF8="1", PYTHONIOENCODING="utf-8"), capture_output=True)
        self.assertEqual(result.returncode, 0, (result.stdout + result.stderr).decode("utf-8", "replace"))
        return result

    def write_csv(self, name, rows, fields=None):
        path = self.root / "out" / name
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=fields or list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)

    def read_csv(self, name):
        return pd.read_csv(self.root / "out" / name, encoding="utf-8-sig")

    def test_template_rerun_preserves_audit_and_final(self):
        audit = self.root / "out" / "manual_audit.xlsx"
        wb = load_workbook(audit)
        wb.active["D2"] = 1
        wb.active["H2"] = "đã kiểm tay"
        wb.save(audit)
        before = audit.read_bytes()
        coding = self.root / "out" / "coding_sheet.xlsx"
        wb = load_workbook(coding)
        wb.active["G2"] = 3
        wb.save(coding)
        self.run_script("00_make_templates.py")
        self.assertEqual(audit.read_bytes(), before)
        self.assertEqual(load_workbook(coding).active["G2"].value, 3)

    def test_auto_fill_keeps_confirmed_rows_and_definitions(self):
        rows = [dict(bank="mb", year=2025, variable=var, value=80, unit="%", source_file="source.pdf",
                     page=2, confirmed_by=label, note="đã kiểm", definition="original")
                for var, label in [("digital_txn_share", "nam"), ("total_assets", "manual-custom"), ("headcount", "range-checked")]]
        self.write_csv("dti_inputs.csv", rows)
        self.write_csv("statement_pages.csv", [], ["bank", "year", "statement", "page"])
        self.write_csv("financial_candidates.csv", [], ["bank", "year", "page", "variable", "parsed_numbers(first4)"])
        self.write_csv("kpi_candidates.csv", [], ["bank", "year", "variable"])
        self.run_script("_fill_dti_inputs.py")
        actual = self.read_csv("dti_inputs.csv")
        self.assertEqual(actual["value"].tolist(), [80, 80, 80])
        self.assertEqual(actual["definition"].tolist(), ["original"] * 3)
        self.assertEqual(actual["confirmed_by"].tolist(), ["nam", "manual-custom", "range-checked"])

    def test_bundle_excludes_unverified_and_uses_latest_report(self):
        rows = [dict(bank="mb", year=2025, variable=var, value=80, unit="%", source_file="source.pdf",
                     page=2, confirmed_by=label, note="đã kiểm", definition="all_txn")
                for var, label in [("digital_txn_share", "nam"), ("total_assets", "range-checked"), ("headcount", "fin-cand")]]
        self.write_csv("dti_inputs.csv", rows)
        cols = ["file", "bank", "year", "doc_type", "pages", "pages_ocr", "avg_chars", "avg_junk", "engines_used", "flag", "lang", "bad_pages"]
        reports = [{key: "" for key in cols} for _ in range(2)]
        for r, flag in zip(reports, ["NGHI LỖI", "OK"]):
            r.update(bank="mb", year=2025, doc_type="bctn", flag=flag)
        self.write_csv("extraction_report.csv", reports, cols)
        (self.root / "config" / "sources.csv").write_text("bank,year,doc_type,url,notes\nmb,2025,bctn,https://example.com/report.pdf,\n", encoding="utf-8-sig")
        self.run_script("12_appendix.py", "--bundle")
        confirmed = self.read_csv("dataset_clean/dti_inputs_confirmed.csv")
        self.assertEqual(confirmed["variable"].tolist(), ["digital_txn_share"])
        self.assertEqual(len(self.read_csv("dataset_clean/dti_inputs_full.csv")), 3)
        dictionary = (self.root / "out/dataset_clean/DATA_DICTIONARY.md").read_text(encoding="utf-8-sig")
        self.assertIn("Có 1 dòng số liệu đã xác nhận trên 3 dòng", dictionary)
        matrix = self.read_csv("availability_matrix.csv")
        mb = matrix[(matrix.bank == "mb") & (matrix.year == 2025)].iloc[0]
        self.assertEqual(mb.bctn, "OK")

    def test_prefill_preserves_source_file(self):
        self.write_csv("financial_candidates.csv", [dict(bank="mb", year=2025, variable="total_assets", page=7,
                       source_file="data/raw/mb/2025_bctc.pdf", unit_hint="VND million",
                       **{"parsed_numbers(first4)": "1,234,567 | 1,000,000", "raw_line (+dòng kế)": "Total assets"})])
        self.run_script("12_appendix.py", "--prefill")
        suggested = self.read_csv("dti_prefill.csv").iloc[0]
        self.assertEqual(suggested.source_file, "data/raw/mb/2025_bctc.pdf")
        self.assertEqual(suggested.value_suggested, 1234567)

    def test_language_split_and_idf_independence(self):
        text_dir = self.root / "data" / "text"
        vi = "Ngân hàng triển khai chuyển đổi số và ngân hàng số với các khách hàng trong năm. " * 15
        en = "The bank deployed digital transformation and digital banking for customers in the year. " * 15
        def add(bank, lang, text):
            (text_dir / f"{bank}_2025_{lang}.jsonl").write_text(json.dumps(dict(bank=bank, year="2025", doc_type="bctn", page=1, lang=lang, text=text), ensure_ascii=False) + "\n", encoding="utf-8")
        add("mb", "vi", vi)
        self.run_script("07_text_mining.py", "--expand")
        before = self.read_csv("tm_criterion.csv")
        add("mb", "en", en)
        self.run_script("09c_auto_score.py")
        self.run_script("07_text_mining.py", "--expand", "--compare")
        after = self.read_csv("tm_criterion.csv")
        self.assertEqual(set(after.lang), {"vi", "en"})
        pd.testing.assert_frame_equal(before.reset_index(drop=True), after[after.lang == "vi"].reset_index(drop=True))

    def test_phobert_build_training_excludes_english(self):
        rows = [dict(criterion="K2", bank="mb", year=2025, doc_type="bctn", lang=lang,
                     snippet=text, valid=1) for lang, text in [("vi", "Ngân hàng triển khai mở tài khoản trực tuyến."), ("en", "The bank launched online onboarding.")]]
        self.write_csv("evidence_candidates.csv", rows)
        self.run_script("09_phobert.py", "--build-train")
        self.assertEqual(self.read_csv("phobert_train.csv")["lang"].tolist(), ["vi"])
        self.run_script("09_phobert.py", "--build-train", "--only-lang", "en")
        self.assertTrue(self.read_csv("phobert_train.csv").empty)


if __name__ == "__main__":
    unittest.main()
